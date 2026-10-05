"""Lectura de los archivos del cliente (CSV o Excel) y normalización de columnas, números, fechas y textos.

Las columnas se reconocen por su nombre sin importar mayúsculas, tildes ni espacios, y admiten sinónimos
habituales en exportaciones de sistemas de caja ("código" = sku, "existencia" = stock, "ticket" = número…).
Un archivo `mapeo.json` en la carpeta permite añadir nombres propios de un cliente: {"ventas": {"Nro Doc": "numero"}}.
"""
from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
from dataclasses import dataclass, field
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zoneinfo import ZoneInfo

ZONA = ZoneInfo("America/Havana")

# Tipo de archivo → columna canónica → sinónimos (ya normalizados).
COLUMNAS: dict[str, dict[str, tuple[str, ...]]] = {
    "productos": {
        "sku": ("sku", "codigo", "cod", "codigo_producto", "referencia", "ref"),
        "nombre": ("nombre", "producto", "descripcion", "articulo", "nombre_producto"),
        "categoria": ("categoria", "familia", "departamento", "grupo"),
        "subcategoria": ("subcategoria", "subfamilia", "subgrupo"),
        "unidad": ("unidad", "um", "unidad_medida"),
        "stock_minimo": ("stock_minimo", "minimo", "stock_min", "existencia_minima"),
        "activo": ("activo", "estado", "habilitado"),
        "precio_usd": ("precio_usd", "pvp_usd", "precio_venta_usd"),
        "precio_cup": ("precio_cup", "pvp_cup", "precio_venta_cup"),
        "precio": ("precio", "pvp", "precio_venta"),
        "coste_usd": ("coste_usd", "costo_usd"),
        "coste_cup": ("coste_cup", "costo_cup"),
        "coste": ("coste", "costo", "coste_unitario", "costo_unitario"),
        "moneda": ("moneda", "divisa"),
        "stock": ("stock", "existencia", "existencias", "stock_actual", "cantidad_en_stock", "inventario"),
        "fecha_stock": ("fecha_stock", "fecha_inventario", "fecha_existencia"),
        "alta": ("alta", "fecha_alta", "creado", "creado_en"),
    },
    "ventas": {
        "numero": ("numero", "ticket", "factura", "numero_venta", "id_venta", "venta", "comprobante", "n_ticket",
                   "nro", "no", "num", "documento"),
        "fecha": ("fecha", "fecha_hora", "fecha_venta", "momento"),
        "hora": ("hora", "hora_venta"),
        "sku": ("sku", "codigo", "cod", "codigo_producto", "referencia", "ref"),
        "nombre": ("nombre", "producto", "descripcion", "articulo"),
        "cantidad": ("cantidad", "cant", "unidades", "uds", "qty"),
        "precio": ("precio", "precio_unitario", "pvp", "precio_venta"),
        "importe": ("importe", "total", "total_linea", "subtotal", "monto"),
        "moneda": ("moneda", "divisa"),
        "canal": ("canal", "tienda", "punto_venta", "origen"),
        "estado": ("estado", "situacion"),
        "metodo_pago": ("metodo_pago", "pago", "forma_pago", "metodo", "forma_de_pago"),
        "tasa": ("tasa", "tasa_usd_cup", "tasa_cambio", "cambio"),
        "coste_unitario_usd": ("coste_unitario_usd", "costo_unitario_usd", "coste_usd", "costo_usd"),
        "anulada_en": ("anulada_en", "fecha_anulacion"),
        "motivo_anulacion": ("motivo_anulacion", "motivo"),
    },
    "devoluciones": {
        "numero": ("numero", "ticket", "factura", "numero_venta", "id_venta", "venta", "comprobante"),
        "sku": ("sku", "codigo", "cod", "codigo_producto", "referencia", "ref"),
        "fecha": ("fecha", "fecha_devolucion", "fecha_hora"),
        "cantidad": ("cantidad", "cant", "unidades", "uds"),
        "reembolso": ("reembolso", "importe", "monto", "devuelto"),
        "moneda": ("moneda", "divisa"),
        "motivo": ("motivo", "causa", "observaciones"),
        "reingresa_stock": ("reingresa_stock", "reingresa", "vuelve_al_stock", "a_stock"),
    },
    "compras": {
        "referencia": ("referencia", "numero", "factura", "pedido", "orden", "id_compra", "documento"),
        "fecha": ("fecha", "fecha_compra", "fecha_recepcion", "fecha_hora"),
        "proveedor": ("proveedor", "suministrador", "vendedor"),
        "sku": ("sku", "codigo", "cod", "codigo_producto"),
        "nombre": ("nombre", "producto", "descripcion", "articulo"),
        "cantidad": ("cantidad", "cant", "unidades", "uds"),
        "coste_unitario": ("coste_unitario", "costo_unitario", "coste", "costo", "precio", "precio_unitario"),
        "moneda": ("moneda", "divisa"),
        "estado": ("estado", "situacion"),
    },
    "inventario": {
        "sku": ("sku", "codigo", "cod", "codigo_producto", "referencia"),
        "stock": ("stock", "existencia", "existencias", "cantidad", "conteo", "stock_actual"),
        "fecha": ("fecha", "fecha_conteo", "fecha_hora"),
        "coste_medio_usd": ("coste_medio_usd", "costo_medio_usd", "coste_usd", "costo_usd", "coste", "costo"),
    },
    "precios": {
        "sku": ("sku", "codigo", "cod", "codigo_producto"),
        "precio": ("precio", "pvp", "precio_venta"),
        "precio_usd": ("precio_usd", "pvp_usd"),
        "precio_cup": ("precio_cup", "pvp_cup"),
        "moneda": ("moneda", "divisa"),
        "desde": ("desde", "fecha", "vigente_desde", "fecha_desde"),
    },
    "tasas": {
        "fecha": ("fecha", "dia"),
        "usd_cup": ("usd_cup", "tasa", "usd", "valor", "tasa_usd_cup"),
    },
    "categorias": {
        "categoria": ("categoria", "familia", "nombre"),
        "linea": ("linea",),
        "plazo_dias": ("plazo_dias", "plazo", "plazo_reposicion", "dias_reposicion"),
        "dias_stock_minimo": ("dias_stock_minimo", "dias_minimo", "cobertura_minima"),
        "basico": ("basico", "basico_canasta", "esencial"),
    },
}

OBLIGATORIAS = {
    "productos": ("sku", "nombre", "categoria"),
    "ventas": ("fecha", "sku", "cantidad"),
    "devoluciones": ("numero", "sku", "fecha", "cantidad"),
    "compras": ("fecha", "sku", "cantidad", "coste_unitario"),
    "inventario": ("sku", "stock"),
    "precios": ("sku", "desde"),
    "tasas": ("fecha", "usd_cup"),
    "categorias": ("categoria",),
}
TIPOS = tuple(COLUMNAS)


class ErrorDato(ValueError):
    """Valor de una celda que no se puede interpretar."""


def normalizar(texto: str) -> str:
    t = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode().strip().lower()
    t = re.sub(r"[^a-z0-9]+", "_", t).strip("_")
    return t


@dataclass
class Fila:
    archivo: str
    fila: int                      # número de fila tal como lo ve el usuario (la cabecera es la 1)
    valores: dict[str, object]

    def get(self, clave: str):
        v = self.valores.get(clave)
        if isinstance(v, str):
            v = v.strip()
            return v or None
        return v

    @property
    def donde(self) -> str:
        return f"{self.archivo}, fila {self.fila}"


@dataclass
class Tabla:
    tipo: str
    archivo: str
    filas: list[Fila] = field(default_factory=list)
    ignoradas: list[str] = field(default_factory=list)   # columnas que no se reconocen


# Archivos -------------------------------------------------------------------------------------------------

def tipo_de(nombre: str) -> str | None:
    """`ventas_2026-10.csv` → ventas. El tipo es el principio del nombre del archivo (o de la hoja de Excel)."""
    n = normalizar(Path(nombre).stem)
    for t in sorted(TIPOS, key=len, reverse=True):
        if n == t or n.startswith(t + "_") or n.startswith(t):
            return t
    return None


def _celdas_csv(ruta: Path) -> list[list[object]]:
    crudo = ruta.read_bytes()
    for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            texto = crudo.decode(codificacion)
            break
        except UnicodeDecodeError:
            continue
    # El separador se deduce de la cabecera, que no lleva decimales: "precio;coste" frente a "1,20;0,85".
    cabecera = next((l for l in texto.splitlines() if l.strip()), "")
    sep = max(";,\t|", key=cabecera.count)
    return [fila for fila in csv.reader(io.StringIO(texto), delimiter=sep)]


def _hojas_xlsx(ruta: Path) -> list[tuple[str, list[list[object]]]]:
    from openpyxl import load_workbook

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        return [(h.title, [list(f) for f in h.iter_rows(values_only=True)]) for h in libro.worksheets]
    finally:
        libro.close()


def _tabla(tipo: str, archivo: str, celdas: list[list[object]], mapeo: dict[str, str]) -> Tabla:
    tabla = Tabla(tipo, archivo)
    # La cabecera es la primera fila con al menos dos celdas no vacías.
    inicio = next((i for i, f in enumerate(celdas) if sum(1 for c in f if c not in (None, "")) >= 2), None)
    if inicio is None:
        return tabla
    sinonimos = {s: canon for canon, alias in COLUMNAS[tipo].items() for s in alias}
    sinonimos.update({normalizar(k): v for k, v in mapeo.items()})
    columnas: list[str | None] = []
    for c in celdas[inicio]:
        n = normalizar(c) if c not in (None, "") else ""
        canon = sinonimos.get(n)
        if canon in columnas:            # la primera columna con ese significado manda
            canon = None
        if canon is None and n:
            tabla.ignoradas.append(str(c))
        columnas.append(canon)
    for i, f in enumerate(celdas[inicio + 1:], start=inicio + 2):
        if all(c in (None, "") or (isinstance(c, str) and not c.strip()) for c in f):
            continue
        valores = {col: f[j] for j, col in enumerate(columnas) if col and j < len(f)}
        tabla.filas.append(Fila(archivo, i, valores))
    return tabla


def leer_carpeta(carpeta: Path) -> tuple[list[Tabla], list[str]]:
    """Lee todos los archivos reconocibles de la carpeta, en orden alfabético (los últimos corrigen a los primeros)."""
    avisos: list[str] = []
    mapeo: dict[str, dict[str, str]] = {}
    if (carpeta / "mapeo.json").exists():
        mapeo = json.loads((carpeta / "mapeo.json").read_text(encoding="utf-8"))
    tablas: list[Tabla] = []
    for ruta in sorted(p for p in carpeta.iterdir() if p.is_file() and not p.name.startswith((".", "~$"))):
        ext = ruta.suffix.lower()
        if ext not in (".csv", ".txt", ".xlsx", ".xlsm") or ruta.name == "mapeo.json":
            if ruta.name not in ("mapeo.json", "LEEME.md", "README.md"):
                avisos.append(f"{ruta.name}: no es CSV ni Excel; se ignora.")
            continue
        if ext in (".xlsx", ".xlsm"):
            tipo_archivo = tipo_de(ruta.name)
            for hoja, celdas in _hojas_xlsx(ruta):
                tipo = tipo_de(hoja) or (tipo_archivo if len(celdas) else None)
                if tipo is None:
                    avisos.append(f"{ruta.name} › {hoja}: el nombre de la hoja no indica qué contiene; se ignora.")
                    continue
                tablas.append(_tabla(tipo, f"{ruta.name} › {hoja}", celdas, mapeo.get(tipo, {})))
        else:
            tipo = tipo_de(ruta.name)
            if tipo is None:
                avisos.append(f"{ruta.name}: el nombre no indica qué contiene (productos, ventas, compras…); se ignora.")
                continue
            tablas.append(_tabla(tipo, ruta.name, _celdas_csv(ruta), mapeo.get(tipo, {})))
    for t in tablas:
        faltan = [c for c in OBLIGATORIAS[t.tipo] if not any(c in f.valores for f in t.filas[:1])]
        if t.filas and faltan:
            avisos.append(f"{t.archivo}: faltan las columnas {', '.join(faltan)}.")
        if t.ignoradas:
            avisos.append(f"{t.archivo}: columnas no reconocidas (se ignoran): {', '.join(t.ignoradas)}.")
    return tablas, avisos


# Valores --------------------------------------------------------------------------------------------------

_MILES = re.compile(r"^-?\d{1,3}([.,]\d{3})+$")


def numero(v: object, nombre: str = "número") -> Decimal | None:
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, bool):
        raise ErrorDato(f"{nombre} no válido: {v!r}")
    if isinstance(v, (int, float, Decimal)):
        return Decimal(str(v))
    s = re.sub(r"(?i)\s|\$|usd|cup|mlc|€", "", str(v))
    if not s:
        return None
    if "," in s and "." in s:                       # el último separador es el decimal: 1.234,56 · 1,234.56
        s = s.replace(".", "").replace(",", ".") if s.rfind(",") > s.rfind(".") else s.replace(",", "")
    elif _MILES.match(s) and not re.match(r"^-?0[.,]", s):
        # 1.250 · 1,250 · 12.345.678: separador de miles (los importes no llevan tres decimales)
        s = s.replace(",", "").replace(".", "")
    else:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation:
        raise ErrorDato(f"{nombre} no válido: {v!r}") from None


_FECHA_ISO = re.compile(r"^(\d{4})-(\d{1,2})-(\d{1,2})(?:[ T](\d{1,2}):(\d{2})(?::(\d{2}))?(?:\.\d+)?)?\s*([+-]\d{2}:?\d{2}|Z)?$")
_FECHA_DMA = re.compile(r"^(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})(?:[ T,]+(\d{1,2}):(\d{2})(?::(\d{2}))?\s*([aApP]\.?\s*[mM]\.?)?)?$")


def fecha_hora(v: object, hora: object = None, nombre: str = "fecha") -> tuple[datetime, bool] | None:
    """Devuelve (momento con zona, solo_fecha). Sin zona se entiende hora de La Habana; los días van primero (dd/mm/aaaa)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    solo_fecha = False
    if isinstance(v, datetime):
        dt = v
        solo_fecha = dt.time() == time(0) and hora is None
    elif isinstance(v, date):
        dt, solo_fecha = datetime.combine(v, time(0)), True
    elif isinstance(v, (int, float)) and 20000 < float(v) < 80000:      # número de serie de Excel
        from datetime import timedelta
        dt = datetime(1899, 12, 30) + timedelta(days=float(v))
        solo_fecha = float(v).is_integer()
    else:
        s = str(v).strip()
        m = _FECHA_ISO.match(s)
        if m:
            a, me, d, h, mi, se, z = m.groups()
            if z:
                try:
                    dt = datetime.fromisoformat(s.replace(" ", "T").replace("Z", "+00:00"))
                except ValueError:
                    raise ErrorDato(f"{nombre} no válida: {v!r}") from None
            else:
                dt = datetime(int(a), int(me), int(d), int(h or 0), int(mi or 0), int(se or 0))
            solo_fecha = h is None
        else:
            m = _FECHA_DMA.match(s)
            if not m:
                raise ErrorDato(f"{nombre} no válida: {v!r} (usa dd/mm/aaaa o aaaa-mm-dd)")
            d, me, a, h, mi, se, ampm = m.groups()
            a = int(a) + (2000 if len(a) == 2 else 0)
            hh = int(h or 0)
            if ampm and ampm.lower().startswith("p") and hh < 12:
                hh += 12
            if ampm and ampm.lower().startswith("a") and hh == 12:
                hh = 0
            try:
                dt = datetime(a, int(me), int(d), hh, int(mi or 0), int(se or 0))
            except ValueError:
                raise ErrorDato(f"{nombre} no válida: {v!r} (usa dd/mm/aaaa o aaaa-mm-dd)") from None
            solo_fecha = h is None
    if hora is not None and solo_fecha:
        h = _hora(hora)
        dt = datetime.combine(dt.date(), h)
        solo_fecha = False
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZONA)
    return dt, solo_fecha


def _hora(v: object) -> time:
    if isinstance(v, time):
        return v
    if isinstance(v, datetime):
        return v.time()
    if isinstance(v, (int, float)) and 0 <= float(v) < 1:               # fracción de día de Excel
        s = round(float(v) * 86400)
        return time(s // 3600, s % 3600 // 60, s % 60)
    m = re.match(r"^(\d{1,2}):(\d{2})(?::(\d{2}))?\s*([aApP]\.?\s*[mM]\.?)?$", str(v).strip())
    if not m:
        raise ErrorDato(f"hora no válida: {v!r}")
    h = int(m.group(1))
    if m.group(4) and m.group(4).lower().startswith("p") and h < 12:
        h += 12
    if m.group(4) and m.group(4).lower().startswith("a") and h == 12:
        h = 0
    return time(h, int(m.group(2)), int(m.group(3) or 0))


SI = {"1", "si", "s", "true", "verdadero", "x", "yes", "y", "activo", "activa", "alta", "habilitado"}
NO = {"0", "no", "n", "false", "falso", "inactivo", "inactiva", "baja", "descatalogado", "deshabilitado"}


def booleano(v: object, defecto: bool, nombre: str = "valor") -> bool:
    if v is None or (isinstance(v, str) and not v.strip()):
        return defecto
    if isinstance(v, bool):
        return v
    n = normalizar(v)
    if n in SI:
        return True
    if n in NO:
        return False
    raise ErrorDato(f"{nombre} no válido: {v!r} (usa sí/no)")


def texto(v: object) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = re.sub(r"\s+", " ", str(v)).strip()
    return s or None


def codigo(v: object) -> str | None:
    """SKU / número de ticket: texto sin espacios sobrantes, en mayúsculas (los SKU de Excel pueden llegar como 1001.0)."""
    s = texto(v)
    return s.upper() if s else None


def opcion(v: object, opciones: dict[str, tuple[str, ...]], defecto: str, nombre: str) -> str:
    if v is None or (isinstance(v, str) and not v.strip()):
        return defecto
    n = normalizar(v)
    for valor, sinonimos in opciones.items():
        if n == valor or n in sinonimos or any(n.startswith(s) for s in sinonimos if len(s) >= 4):
            return valor
    raise ErrorDato(f"{nombre} no válido: {v!r} (valores: {', '.join(opciones)})")


MONEDAS = {"USD": ("usd", "dolar", "dolares", "us", "divisa"), "CUP": ("cup", "peso", "pesos", "mn", "moneda_nacional", "cuc")}
CANALES = {"tienda_fisica": ("tienda", "fisica", "local", "presencial", "mostrador", "tienda_fisica"),
           "web": ("web", "online", "internet", "domicilio", "envio", "tienda_online", "whatsapp")}
ESTADOS_VENTA = {"completada": ("completada", "completado", "pagada", "pagado", "cobrada", "ok", "vendida", "cerrada"),
                 "anulada": ("anulada", "anulado", "cancelada", "cancelado", "void", "nula")}
METODOS = {"efectivo": ("efectivo", "cash", "contado"),
           "transferencia": ("transferencia", "transfermovil", "enzona", "tarjeta", "qr", "transf")}
ESTADOS_COMPRA = {"recibida": ("recibida", "recibido", "entregada", "entregado", "completada", "cerrada"),
                  "pendiente": ("pendiente", "en_camino", "pedida", "pedido", "abierta"),
                  "cancelada": ("cancelada", "cancelado", "anulada", "anulado")}
LINEAS = {"Mercado": ("mercado",), "Envíos": ("envios", "envio")}


def moneda(v: object) -> str:
    return opcion(v, MONEDAS, "USD", "moneda")
